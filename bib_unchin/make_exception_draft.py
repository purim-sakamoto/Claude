"""過去の運賃チェック結果から「BIB運賃付け合わせ」用の例外リスト素案を作る。

入力: 運賃チェックシート(xlsm) の「張り付けシート」（マクロ実行後の状態）
出力: 例外リスト_素案.xlsx
"""
import collections
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src, dst = sys.argv[1], sys.argv[2]

rows = list(openpyxl.load_workbook(src, read_only=True, data_only=True)["張り付けシート"].iter_rows(values_only=True))[1:]
# 列: 0売上日 1売上番号 2伝票種類 3受注番号 4出荷日 5納入先ｺｰﾄﾞ 6納入先名 8住所1 12運送会社名 15県名 16想定単価 17重量 19売上金額

orders = collections.defaultdict(list)
for r in rows:
    orders[r[3]].append(r)


def num(v):
    return v if isinstance(v, (int, float)) else None


detail = []  # 受注単位の NG
per_code = collections.defaultdict(lambda: {"orders": 0, "small_ok": 0, "ng": []})
for no, lines in orders.items():
    prod = [x for x in lines if (num(x[17]) or 0) > 0]
    frt = [x for x in lines if x not in prod]
    head = prod[0] if prod else lines[0]
    code = head[5]
    info = per_code[code]
    info.update(name=head[6].strip(), addr=head[8], carrier=head[12], pref=head[15])
    info["orders"] += 1
    if len(lines) == 1:
        rec = dict(no=no, date=head[4], kind="運賃行なし", weight=None, unit=num(head[16]), exp=None, act=None)
    else:
        p, f = prod[0], frt[0]
        unit, weight, act = num(p[16]), p[17], f[19]
        exp = round(unit * weight) if unit is not None else None
        if exp == act:
            if weight <= 80:
                info["small_ok"] += 1
            continue
        if unit is None:
            kind = "単価表に県なし"
        elif p[12] == "博運社":
            kind = "離島（博運社）"
        elif weight <= 80:
            kind = "小口（80kg以下）"
        else:
            kind = "単価相違"
        rec = dict(no=no, date=p[4], kind=kind, weight=weight, unit=unit, exp=exp, act=act)
    rec.update(code=code, name=info["name"], pref=info["pref"], carrier=info["carrier"])
    info["ng"].append(rec)
    detail.append(rec)

HANDLING = {
    "運賃行なし": "運賃行なしでもOK（引取・合積み等）",
    "離島（博運社）": "単価チェック対象外（離島運賃）",
    "小口（80kg以下）": "単価チェック対象外（小口運賃）",
    "単価表に県なし": "単価チェック対象外（単価表に県なし）",
    "単価相違": "要確認",
}

exc = []
for code, info in per_code.items():
    if not info["ng"]:
        continue
    kinds = collections.Counter(r["kind"] for r in info["ng"])
    kind = kinds.most_common(1)[0][0]
    n_ng = len(info["ng"])
    notes = []
    if kind == "運賃行なし" and n_ng < info["orders"]:
        notes.append(f"全{info['orders']}受注のうち{n_ng}件だけ運賃なし → 納入先単位より受注単位の例外が妥当かも")
    if kind == "小口（80kg以下）" and info["small_ok"]:
        notes.append(f"80kg以下でも単価どおりの回が{info['small_ok']}件あり")
    if len(kinds) > 1:
        notes.append("区分混在: " + "、".join(f"{k}{v}件" for k, v in kinds.items()))
    exc.append(dict(code=code, name=info["name"], addr=info["addr"], pref=info["pref"], carrier=info["carrier"],
                    kind=kind, handling=HANDLING[kind], ng=n_ng, orders=info["orders"],
                    refs="、".join(r["no"] for r in info["ng"]), note=" / ".join(notes)))

order = list(HANDLING)
exc.sort(key=lambda e: (order.index(e["kind"]), e["pref"] or "", e["code"]))
detail.sort(key=lambda d: (order.index(d["kind"]), d["code"], d["date"]))

wb = openpyxl.Workbook()
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
KIND_FILL = {
    "運賃行なし": "FCE4D6", "離島（博運社）": "DDEBF7", "小口（80kg以下）": "E2EFDA",
    "単価表に県なし": "FFF2CC", "単価相違": "F8CBAD",
}


def write(ws, header, data, widths, kind_col=None):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in data:
        ws.append(row)
        if kind_col is not None:
            fill = PatternFill("solid", fgColor=KIND_FILL[row[kind_col]])
            ws.cell(ws.max_row, kind_col + 1).fill = fill
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


ws = wb.active
ws.title = "例外リスト"
write(ws,
      ["納入先コード", "納入先名", "住所", "県", "運送会社", "例外区分", "照合時の扱い（案）",
       "該当件数", "期間中の受注数", "根拠の受注番号", "気になる点", "確認（○/×）", "確認メモ"],
      [[e["code"], e["name"], e["addr"], e["pref"], e["carrier"], e["kind"], e["handling"],
        e["ng"], e["orders"], e["refs"], e["note"], "", ""] for e in exc],
      [13, 34, 38, 8, 12, 16, 30, 9, 9, 30, 44, 10, 30], kind_col=5)

ws = wb.create_sheet("NG明細")
write(ws,
      ["受注番号", "出荷日", "納入先コード", "納入先名", "県", "運送会社", "例外区分",
       "重量(kg)", "想定単価", "想定運賃", "請求運賃", "差額", "実質単価"],
      [[d["no"], d["date"], d["code"], d["name"], d["pref"], d["carrier"], d["kind"], d["weight"], d["unit"],
        d["exp"], d["act"],
        (d["act"] - d["exp"]) if d["exp"] is not None and d["act"] is not None else None,
        round(d["act"] / d["weight"], 2) if d["act"] and d["weight"] else None] for d in detail],
      [12, 11, 13, 34, 8, 12, 16, 9, 9, 10, 10, 10, 9], kind_col=6)

ws = wb.create_sheet("集計と読み方")
kc = collections.Counter(d["kind"] for d in detail)
lines = [
    ["BIB運賃付け合わせ 例外リスト（素案）"],
    [],
    ["元データ", "運賃チェックシート ver.0.02（2026/04/01〜04/27、在庫伝票、引取・自社便は除外済み）"],
    ["受注数", len(orders)],
    ["突合方法", "受注番号ごとに「商品行」と「運賃行」を1対1で組み、想定運賃（県別単価×重量）と請求運賃を比較"],
    ["結果", f"一致 {len(orders) - len(detail)} 件 / 不一致 {len(detail)} 件（単価NGの件数は既存シートの×と同じ）"],
    [],
    ["例外区分", "件数", "照合時の扱い（案）", "所見"],
    ["運賃行なし", kc["運賃行なし"], HANDLING["運賃行なし"], "同じ納入先で毎回運賃なし（例: 住友金属鉱山 菱刈鉱山）は納入先単位で登録。たまにだけのものは受注単位の例外が妥当"],
    ["離島（博運社）", kc["離島（博運社）"], HANDLING["離島（博運社）"], "奄美・種子島・徳之島・対馬などで、実質 28〜40円/kg。運送会社＝博運社をまとめて除外にするか、離島用の単価を持つかを決めたい"],
    ["小口（80kg以下）", kc["小口（80kg以下）"], HANDLING["小口（80kg以下）"], "1,380〜3,000円の固定額で請求されている。一方で80kg以下でも単価どおりの回が44件ある。納入先ごとに決まっているのか、別のルールがあるのかを確認したい"],
    ["単価表に県なし", kc["単価表に県なし"], HANDLING["単価表に県なし"], "沖縄県（トーホー沖縄 石垣）。単価表に追加するか例外にする"],
    ["単価相違", kc["単価相違"], HANDLING["単価相違"], "上のどれにも当てはまらない本当の単価違い（今回はなし）"],
    [],
    ["使い方", "「例外リスト」の確認欄に○/×を付けてください。○の行が「BIB運賃付け合わせ」の例外マスタになります。"],
]
for l in lines:
    ws.append(l)
ws["A1"].font = Font(bold=True, size=14)
for c in ws[8]:
    c.fill, c.font = HEAD_FILL, HEAD_FONT
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 34
ws.column_dimensions["D"].width = 90
wb.move_sheet("集計と読み方", offset=-2)
wb.save(dst)
print(f"例外 {len(exc)} 納入先 / 明細 {len(detail)} 件", dict(kc))
